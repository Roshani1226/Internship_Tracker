import tkinter as tk
from tkinter import ttk, messagebox
import mysql.connector
from openpyxl import Workbook, load_workbook
import os
from datetime import datetime


DB_HOST = "localhost"
DB_USER = "root"
DB_PASSWORD = "roshu1221"     
DB_NAME = "internship_tracker"


EXCEL_FILE = "internship_data.xlsx"


def create_database():
    try:
        conn = mysql.connector.connect(
            host=DB_HOST,
            user=DB_USER,
            password=DB_PASSWORD
        )

        cursor = conn.cursor()

        cursor.execute(
            f"CREATE DATABASE IF NOT EXISTS {DB_NAME}"
        )

        cursor.close()
        conn.close()

    except mysql.connector.Error as e:
        messagebox.showerror(
            "Database Error",
            f"Could not create database.\n\n{e}"
        )


def connect_database():
    return mysql.connector.connect(
        host=DB_HOST,
        user=DB_USER,
        password=DB_PASSWORD,
        database=DB_NAME
    )


def create_table():
    try:
        conn = connect_database()
        cursor = conn.cursor()

        query = """
        CREATE TABLE IF NOT EXISTS applications (
            id INT AUTO_INCREMENT PRIMARY KEY,
            company_name VARCHAR(100) NOT NULL,
            internship_role VARCHAR(100) NOT NULL,
            location VARCHAR(100),
            internship_type VARCHAR(30),
            application_date DATE,
            deadline DATE,
            status VARCHAR(30),
            interview_date DATE,
            stipend VARCHAR(50),
            duration VARCHAR(50),
            website VARCHAR(255),
            notes TEXT
        )
        """

        cursor.execute(query)

        conn.commit()
        cursor.close()
        conn.close()

    except mysql.connector.Error as e:
        messagebox.showerror(
            "Database Error",
            f"Could not create table.\n\n{e}"
        )



def create_excel_file():

    if not os.path.exists(EXCEL_FILE):

        workbook = Workbook()
        sheet = workbook.active

        sheet.title = "Internship Applications"

        headers = [
            "ID",
            "Company Name",
            "Internship Role",
            "Location",
            "Internship Type",
            "Application Date",
            "Deadline",
            "Status",
            "Interview Date",
            "Stipend",
            "Duration",
            "Website",
            "Notes"
        ]

        sheet.append(headers)

        workbook.save(EXCEL_FILE)


def sync_mysql_to_excel():

    try:
        conn = connect_database()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT
                id,
                company_name,
                internship_role,
                location,
                internship_type,
                application_date,
                deadline,
                status,
                interview_date,
                stipend,
                duration,
                website,
                notes
            FROM applications
            ORDER BY id
        """)

        records = cursor.fetchall()

        cursor.close()
        conn.close()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Internship Applications"

        headers = [
            "ID",
            "Company Name",
            "Internship Role",
            "Location",
            "Internship Type",
            "Application Date",
            "Deadline",
            "Status",
            "Interview Date",
            "Stipend",
            "Duration",
            "Website",
            "Notes"
        ]

        sheet.append(headers)

        for record in records:
            sheet.append(list(record))

        # Make columns wider
        widths = [8, 20, 25, 18, 18, 18, 18,
                  18, 18, 15, 15, 30, 30]

        for i, width in enumerate(widths, start=1):
            sheet.column_dimensions[
                chr(64 + i)
            ].width = width

        workbook.save(EXCEL_FILE)

    except Exception as e:
        messagebox.showerror(
            "Excel Error",
            f"Could not update Excel file.\n\n{e}"
        )



def add_application():

    if company_var.get().strip() == "":
        messagebox.showwarning(
            "Missing Information",
            "Please enter Company Name."
        )
        return

    if role_var.get().strip() == "":
        messagebox.showwarning(
            "Missing Information",
            "Please enter Internship Role."
        )
        return

    try:

        conn = connect_database()
        cursor = conn.cursor()

        query = """
        INSERT INTO applications
        (
            company_name,
            internship_role,
            location,
            internship_type,
            application_date,
            deadline,
            status,
            interview_date,
            stipend,
            duration,
            website,
            notes
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """

        values = (
            company_var.get(),
            role_var.get(),
            location_var.get(),
            type_var.get(),
            application_date_var.get(),
            deadline_var.get(),
            status_var.get(),
            interview_date_var.get(),
            stipend_var.get(),
            duration_var.get(),
            website_var.get(),
            notes_text.get("1.0", tk.END).strip()
        )

        cursor.execute(query, values)

        conn.commit()

        cursor.close()
        conn.close()

        # Update Excel
        sync_mysql_to_excel()

        load_data()
        update_dashboard()
        clear_fields()

        messagebox.showinfo(
            "Success",
            "Internship application saved successfully!\n\n"
            "Data saved in:\n"
            "✓ MySQL Database\n"
            "✓ Excel File"
        )

    except mysql.connector.Error as e:

        messagebox.showerror(
            "Database Error",
            str(e)
        )



def load_data():

    for item in tree.get_children():
        tree.delete(item)

    try:

        conn = connect_database()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT
                id,
                company_name,
                internship_role,
                location,
                internship_type,
                application_date,
                deadline,
                status,
                interview_date,
                stipend,
                duration,
                website,
                notes
            FROM applications
            ORDER BY id DESC
        """)

        records = cursor.fetchall()

        for record in records:
            tree.insert(
                "",
                tk.END,
                values=record
            )

        cursor.close()
        conn.close()

        update_dashboard()

    except mysql.connector.Error as e:

        messagebox.showerror(
            "Database Error",
            str(e)
        )


def search_application():

    search_value = search_var.get().strip()

    if search_value == "":
        load_data()
        return

    for item in tree.get_children():
        tree.delete(item)

    try:

        conn = connect_database()
        cursor = conn.cursor()

        query = """
        SELECT
            id,
            company_name,
            internship_role,
            location,
            internship_type,
            application_date,
            deadline,
            status,
            interview_date,
            stipend,
            duration,
            website,
            notes
        FROM applications
        WHERE company_name LIKE %s
        OR internship_role LIKE %s
        OR status LIKE %s
        ORDER BY id DESC
        """

        value = "%" + search_value + "%"

        cursor.execute(
            query,
            (value, value, value)
        )

        records = cursor.fetchall()

        for record in records:
            tree.insert(
                "",
                tk.END,
                values=record
            )

        cursor.close()
        conn.close()

    except mysql.connector.Error as e:

        messagebox.showerror(
            "Search Error",
            str(e)
        )


def select_application(event):

    selected = tree.focus()

    if not selected:
        return

    values = tree.item(selected, "values")

    if not values:
        return

    company_var.set(values[1])
    role_var.set(values[2])
    location_var.set(values[3])
    type_var.set(values[4])
    application_date_var.set(values[5])
    deadline_var.set(values[6])
    status_var.set(values[7])
    interview_date_var.set(values[8])
    stipend_var.set(values[9])
    duration_var.set(values[10])
    website_var.set(values[11])

    notes_text.delete("1.0", tk.END)

    if len(values) > 12:
        notes_text.insert("1.0", values[12])



def update_application():

    selected = tree.focus()

    if not selected:
        messagebox.showwarning(
            "Select Record",
            "Please select an application to update."
        )
        return

    values = tree.item(selected, "values")

    application_id = values[0]

    try:

        conn = connect_database()
        cursor = conn.cursor()

        query = """
        UPDATE applications
        SET
            company_name=%s,
            internship_role=%s,
            location=%s,
            internship_type=%s,
            application_date=%s,
            deadline=%s,
            status=%s,
            interview_date=%s,
            stipend=%s,
            duration=%s,
            website=%s,
            notes=%s
        WHERE id=%s
        """

        data = (
            company_var.get(),
            role_var.get(),
            location_var.get(),
            type_var.get(),
            application_date_var.get(),
            deadline_var.get(),
            status_var.get(),
            interview_date_var.get(),
            stipend_var.get(),
            duration_var.get(),
            website_var.get(),
            notes_text.get("1.0", tk.END).strip(),
            application_id
        )

        cursor.execute(query, data)

        conn.commit()

        cursor.close()
        conn.close()

        # Update Excel
        sync_mysql_to_excel()

        load_data()
        clear_fields()

        messagebox.showinfo(
            "Updated",
            "Application updated successfully!\n\n"
            "MySQL and Excel have been updated."
        )

    except mysql.connector.Error as e:

        messagebox.showerror(
            "Update Error",
            str(e)
        )



def delete_application():

    selected = tree.focus()

    if not selected:
        messagebox.showwarning(
            "Select Record",
            "Please select an application to delete."
        )
        return

    values = tree.item(selected, "values")

    application_id = values[0]

    confirm = messagebox.askyesno(
        "Confirm Delete",
        "Are you sure you want to delete this application?"
    )

    if not confirm:
        return

    try:

        conn = connect_database()
        cursor = conn.cursor()

        cursor.execute(
            "DELETE FROM applications WHERE id=%s",
            (application_id,)
        )

        conn.commit()

        cursor.close()
        conn.close()

        # Update Excel
        sync_mysql_to_excel()

        load_data()
        clear_fields()

        messagebox.showinfo(
            "Deleted",
            "Application deleted successfully!\n\n"
            "MySQL and Excel have been updated."
        )

    except mysql.connector.Error as e:

        messagebox.showerror(
            "Delete Error",
            str(e)
        )


def clear_fields():

    company_var.set("")
    role_var.set("")
    location_var.set("")
    type_var.set("Online")
    application_date_var.set("")
    deadline_var.set("")
    status_var.set("Applied")
    interview_date_var.set("")
    stipend_var.set("")
    duration_var.set("")
    website_var.set("")

    notes_text.delete(
        "1.0",
        tk.END
    )

    for item in tree.selection():
        tree.selection_remove(item)


def update_dashboard():

    try:

        conn = connect_database()
        cursor = conn.cursor()

        cursor.execute(
            "SELECT COUNT(*) FROM applications"
        )

        total = cursor.fetchone()[0]

        cursor.execute(
            "SELECT COUNT(*) FROM applications WHERE status='Applied'"
        )

        applied = cursor.fetchone()[0]

        cursor.execute(
            "SELECT COUNT(*) FROM applications WHERE status='Shortlisted'"
        )

        shortlisted = cursor.fetchone()[0]

        cursor.execute(
            "SELECT COUNT(*) FROM applications WHERE status='Selected'"
        )

        selected = cursor.fetchone()[0]

        cursor.close()
        conn.close()

        total_label.config(
            text=f"Total Applications\n{total}"
        )

        applied_label.config(
            text=f"Applied\n{applied}"
        )

        shortlisted_label.config(
            text=f"Shortlisted\n{shortlisted}"
        )

        selected_label.config(
            text=f"Selected\n{selected}"
        )

    except mysql.connector.Error:
        pass


root = tk.Tk()

root.title(
    "Digital Internship Application Tracker"
)

root.geometry(
    "1450x800"
)

root.configure(
    bg="#f4f6f8"
)



title_label = tk.Label(
    root,
    text="Digital Internship Application Tracker",
    font=("Arial", 24, "bold"),
    bg="#f4f6f8"
)

title_label.pack(
    pady=15
)



dashboard = tk.Frame(
    root,
    bg="#f4f6f8"
)

dashboard.pack(
    pady=5
)


total_label = tk.Label(
    dashboard,
    text="Total Applications\n0",
    font=("Arial", 14, "bold"),
    width=20,
    height=3,
    relief="ridge"
)

total_label.grid(
    row=0,
    column=0,
    padx=10
)


applied_label = tk.Label(
    dashboard,
    text="Applied\n0",
    font=("Arial", 14, "bold"),
    width=20,
    height=3,
    relief="ridge"
)

applied_label.grid(
    row=0,
    column=1,
    padx=10
)


shortlisted_label = tk.Label(
    dashboard,
    text="Shortlisted\n0",
    font=("Arial", 14, "bold"),
    width=20,
    height=3,
    relief="ridge"
)

shortlisted_label.grid(
    row=0,
    column=2,
    padx=10
)


selected_label = tk.Label(
    dashboard,
    text="Selected\n0",
    font=("Arial", 14, "bold"),
    width=20,
    height=3,
    relief="ridge"
)

selected_label.grid(
    row=0,
    column=3,
    padx=10
)


form_frame = tk.LabelFrame(
    root,
    text="Internship Application Details",
    font=("Arial", 12, "bold"),
    padx=10,
    pady=10
)

form_frame.pack(
    fill="x",
    padx=20,
    pady=10
)


# Variables

company_var = tk.StringVar()
role_var = tk.StringVar()
location_var = tk.StringVar()
type_var = tk.StringVar(value="Online")
application_date_var = tk.StringVar()
deadline_var = tk.StringVar()
status_var = tk.StringVar(value="Applied")
interview_date_var = tk.StringVar()
stipend_var = tk.StringVar()
duration_var = tk.StringVar()
website_var = tk.StringVar()


# Row 1

tk.Label(
    form_frame,
    text="Company Name"
).grid(row=0, column=0, padx=5, pady=5)

tk.Entry(
    form_frame,
    textvariable=company_var,
    width=25
).grid(row=0, column=1, padx=5)


tk.Label(
    form_frame,
    text="Internship Role"
).grid(row=0, column=2, padx=5)

tk.Entry(
    form_frame,
    textvariable=role_var,
    width=25
).grid(row=0, column=3, padx=5)


tk.Label(
    form_frame,
    text="Location"
).grid(row=0, column=4, padx=5)

tk.Entry(
    form_frame,
    textvariable=location_var,
    width=20
).grid(row=0, column=5, padx=5)


# Row 2

tk.Label(
    form_frame,
    text="Internship Type"
).grid(row=1, column=0, padx=5, pady=5)

ttk.Combobox(
    form_frame,
    textvariable=type_var,
    values=["Online", "Offline", "Hybrid"],
    state="readonly",
    width=22
).grid(row=1, column=1, padx=5)


tk.Label(
    form_frame,
    text="Application Date"
).grid(row=1, column=2, padx=5)

tk.Entry(
    form_frame,
    textvariable=application_date_var,
    width=25
).grid(row=1, column=3, padx=5)


tk.Label(
    form_frame,
    text="Deadline"
).grid(row=1, column=4, padx=5)

tk.Entry(
    form_frame,
    textvariable=deadline_var,
    width=20
).grid(row=1, column=5, padx=5)


# Row 3

tk.Label(
    form_frame,
    text="Status"
).grid(row=2, column=0, padx=5, pady=5)

ttk.Combobox(
    form_frame,
    textvariable=status_var,
    values=[
        "Applied",
        "Shortlisted",
        "Interview",
        "Selected",
        "Rejected",
        "Withdrawn"
    ],
    state="readonly",
    width=22
).grid(row=2, column=1, padx=5)


tk.Label(
    form_frame,
    text="Interview Date"
).grid(row=2, column=2, padx=5)

tk.Entry(
    form_frame,
    textvariable=interview_date_var,
    width=25
).grid(row=2, column=3, padx=5)


tk.Label(
    form_frame,
    text="Stipend"
).grid(row=2, column=4, padx=5)

tk.Entry(
    form_frame,
    textvariable=stipend_var,
    width=20
).grid(row=2, column=5, padx=5)


# Row 4

tk.Label(
    form_frame,
    text="Duration"
).grid(row=3, column=0, padx=5, pady=5)

tk.Entry(
    form_frame,
    textvariable=duration_var,
    width=25
).grid(row=3, column=1, padx=5)


tk.Label(
    form_frame,
    text="Website"
).grid(row=3, column=2, padx=5)

tk.Entry(
    form_frame,
    textvariable=website_var,
    width=25
).grid(row=3, column=3, padx=5)


tk.Label(
    form_frame,
    text="Notes"
).grid(row=3, column=4, padx=5)

notes_text = tk.Text(
    form_frame,
    width=25,
    height=3
)

notes_text.grid(
    row=3,
    column=5,
    padx=5
)


button_frame = tk.Frame(
    root,
    bg="#f4f6f8"
)

button_frame.pack(
    pady=10
)


tk.Button(
    button_frame,
    text="Add Application",
    width=18,
    command=add_application
).grid(row=0, column=0, padx=5)


tk.Button(
    button_frame,
    text="Update",
    width=18,
    command=update_application
).grid(row=0, column=1, padx=5)


tk.Button(
    button_frame,
    text="Delete",
    width=18,
    command=delete_application
).grid(row=0, column=2, padx=5)


tk.Button(
    button_frame,
    text="Clear",
    width=18,
    command=clear_fields
).grid(row=0, column=3, padx=5)



search_frame = tk.Frame(
    root,
    bg="#f4f6f8"
)

search_frame.pack(
    pady=5
)


tk.Label(
    search_frame,
    text="Search:"
).pack(
    side="left",
    padx=5
)


search_var = tk.StringVar()

tk.Entry(
    search_frame,
    textvariable=search_var,
    width=40
).pack(
    side="left",
    padx=5
)


tk.Button(
    search_frame,
    text="Search",
    command=search_application
).pack(
    side="left",
    padx=5
)


tk.Button(
    search_frame,
    text="Show All",
    command=load_data
).pack(
    side="left",
    padx=5
)



table_frame = tk.Frame(root)

table_frame.pack(
    fill="both",
    expand=True,
    padx=20,
    pady=10
)


columns = (
    "ID",
    "Company",
    "Role",
    "Location",
    "Type",
    "Application Date",
    "Deadline",
    "Status",
    "Interview Date",
    "Stipend",
    "Duration",
    "Website",
    "Notes"
)


tree = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings"
)


for column in columns:

    tree.heading(
        column,
        text=column
    )

    tree.column(
        column,
        width=120
    )


tree.column("ID", width=50)
tree.column("Company", width=130)
tree.column("Role", width=150)
tree.column("Location", width=100)
tree.column("Type", width=90)
tree.column("Status", width=100)


scrollbar_y = ttk.Scrollbar(
    table_frame,
    orient="vertical",
    command=tree.yview
)

scrollbar_x = ttk.Scrollbar(
    table_frame,
    orient="horizontal",
    command=tree.xview
)


tree.configure(
    yscrollcommand=scrollbar_y.set,
    xscrollcommand=scrollbar_x.set
)


tree.pack(
    side="top",
    fill="both",
    expand=True
)

scrollbar_y.pack(
    side="right",
    fill="y"
)

scrollbar_x.pack(
    side="bottom",
    fill="x"
)


tree.bind(
    "<ButtonRelease-1>",
    select_application
)



create_database()
create_table()
create_excel_file()
sync_mysql_to_excel()
load_data()
root.mainloop()